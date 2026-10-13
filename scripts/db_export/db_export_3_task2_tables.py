"""Populate 4 Task 2 tables: Paragraph/Figure/ClaimElement/GroundingGold. Run db_export_1_core_tables.py and db_export_2_fill_citation_category.py first.
Paragraph/Figure: PRIOR only, because EPO_XML_DIR contains PRIOR full XML but not TARGET; TARGET remains pending and may require extraction from original EPO_SRPT archives.
ClaimElement: split PRIOR XML by claim/claim-text with claim_no. TARGET uses approximate '- ' splitting without known claim_no and is less precise.
GroundingGold: transfer rel_claims/paragraphs_expanded/passage_type_grouped. Leave claim_element_id empty because approximate TARGET elements cannot join reliably; retain target_rel_claims instead. Restrict both query and cited IDs to patent_content to preserve consistency: epo_dataset covers 12 years and is much broader.
Usage: python3 scripts/db_export/db_export_3_task2_tables.py [--citations-dir /path/to/epo_dataset].
"""
import os
import sys
import re
import csv
import json
import glob
import argparse

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from config.settings import DATA_DIR, EPO_XML_DIR, PROJECT_DIR
from data_parser.patent_fulltext_parser import parse_paragraphs, parse_figures, parse_claim_elements

try:
    import openpyxl
except ImportError:
    openpyxl = None


def parse_prior_xml(xml_path):
    """Read one PRIOR full-text XML file and delegate to data_parser.patent_fulltext_parser; parsing is reused for OPS responses with the same ep-patent-document DTD."""
    with open(xml_path, 'r', encoding='utf-8', errors='ignore') as f:
        text = f.read()
    paragraphs = parse_paragraphs(text)
    figures = parse_figures(text)
    claim_elements = parse_claim_elements(text)
    return paragraphs, figures, claim_elements


def split_target_claim_heuristic(text):
    """Approximate TARGET claim elements by '- ' splitting without full XML; claim_no is unknown."""
    if not text:
        return []
    return [p.strip() for p in re.split(r';?\s+-\s+', text) if p.strip()]


def build_paragraph_and_figure_and_claimelement(records):
    para_rows, fig_rows, ce_rows = [], [], []
    seen_prior = set()

    for r in records:
        prior_id, prior_xml = r.get('prior_patent_id'), r.get('prior_xml_file')
        if prior_id and prior_xml and prior_id not in seen_prior:
            seen_prior.add(prior_id)
            xml_path = os.path.join(EPO_XML_DIR, prior_xml)
            if os.path.exists(xml_path):
                paragraphs, figures, claim_elements = parse_prior_xml(xml_path)
                for p in paragraphs:
                    para_rows.append({'para_id': f"{prior_id}_p{p['para_no']}", 'patent_id': prior_id,
                                       'para_no': p['para_no'], 'text': p['text']})
                for fid, meta in figures.items():
                    fig_rows.append({'fig_id': f"{prior_id}_{fid}", 'patent_id': prior_id,
                                      'fig_no': meta['fig_no'], 'caption': meta['caption'],
                                      'image_path': meta['image_path'], 'ocsr_smiles': None})
                for claim_no, element_no, el_text in claim_elements:
                    ce_rows.append({'element_id': f"{prior_id}_c{claim_no}_e{element_no}", 'patent_id': prior_id,
                                     'claim_no': claim_no, 'element_no': element_no,
                                     'element_text': el_text, 'element_smiles': None})

        target_id = r.get('target_patent_id')
        if target_id:
            for i, el_text in enumerate(split_target_claim_heuristic(r.get('target_claim')), start=1):
                ce_rows.append({'element_id': f"{target_id}_TGT_e{i}", 'patent_id': target_id,
                                 'claim_no': None, 'element_no': i,
                                 'element_text': el_text, 'element_smiles': None})

    return para_rows, fig_rows, ce_rows


def build_grounding_gold(citations_dir, known_patent_ids):
    """Retain only rows whose query and cited IDs are both in known_patent_ids. epo_dataset covers 12 years and exceeds this corpus; outside patents would break table consistency. Broaden this filter only when expanding the corpus."""
    if openpyxl is None:
        print('[WARNING] Skipping GroundingGold without openpyxl: pip install openpyxl')
        return []
    rows = []
    skipped_out_of_corpus = 0
    for path in sorted(glob.glob(os.path.join(citations_dir, 'parsed_citations_*.xlsx'))):
        print(f'Reading for GroundingGold: {os.path.basename(path)} ...')
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = next(it)
        idx = {h: i for i, h in enumerate(header)}
        needed = ('ID', 'dnum', 'country', 'rel_claims', 'paragraphs_expanded', 'passage_type_grouped', 'citation_id', 'file_year')
        if not all(c in idx for c in needed):
            wb.close()
            continue
        for row in it:
            if row[idx['country']] != 'EP':
                continue
            target_id, cited_id = row[idx['ID']], row[idx['dnum']]
            if not target_id or not cited_id:
                continue
            if target_id not in known_patent_ids or cited_id not in known_patent_ids:
                skipped_out_of_corpus += 1
                continue
            paragraphs_expanded = row[idx['paragraphs_expanded']]
            evidence_units = [u.strip('[]') for u in str(paragraphs_expanded).split(',') if u.strip()] \
                if paragraphs_expanded else []
            evidence_type = 'figure' if 'figure' in str(row[idx['passage_type_grouped']] or '').lower() \
                else ('paragraph' if evidence_units else 'unknown')
            rows.append({
                'ground_id': f"{row[idx['file_year']]}_{row[idx['citation_id']]}_{target_id}",
                'query_patent_id': target_id,
                'claim_element_id': None,  # Approximate TARGET ClaimElement cannot join reliably; see target_rel_claims.
                'target_rel_claims': row[idx['rel_claims']],
                'cited_patent_id': cited_id,
                'evidence_type': evidence_type,
                'evidence_unit_id': '|'.join(evidence_units) if evidence_units else None,
                'source': 'office_action',
            })
        wb.close()
    print(f'Excluded citation rows with patents outside patent_content: {skipped_out_of_corpus} records')
    return rows


def write_csv(path, fieldnames, rows):
    with open(path, 'w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)
    print(f'{path}: {len(rows)} rows')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--citations-dir', default=os.path.join(os.path.dirname(PROJECT_DIR), 'epo_dataset'))
    args = parser.parse_args()

    dataset_path = os.path.join(DATA_DIR, 'reference', 'epo_section_c_benchmark_dataset.json')
    if not os.path.exists(dataset_path):
        print(f'[FAIL] {dataset_path} is missing. Run python3 run_chem_rag.py first.')
        sys.exit(1)
    with open(dataset_path, 'r', encoding='utf-8') as f:
        records = json.load(f)

    out_dir = os.path.join(DATA_DIR, 'db_export')
    os.makedirs(out_dir, exist_ok=True)

    if not os.path.exists(EPO_XML_DIR):
        print(f'[WARNING] {EPO_XML_DIR} is unavailable; check the external drive connection.')

    # Compute patent IDs actually present in the corpus to keep patent_content.csv
    # consistent and restrict grounding_gold to that set.
    known_patent_ids = set()
    for r in records:
        known_patent_ids.update(pid for pid in (r.get('target_patent_id'), r.get('prior_patent_id')) if pid)

    para_rows, fig_rows, ce_rows = build_paragraph_and_figure_and_claimelement(records)
    write_csv(os.path.join(out_dir, 'paragraph.csv'), ['para_id', 'patent_id', 'para_no', 'text'], para_rows)
    write_csv(os.path.join(out_dir, 'figure.csv'),
              ['fig_id', 'patent_id', 'fig_no', 'caption', 'image_path', 'ocsr_smiles'], fig_rows)
    write_csv(os.path.join(out_dir, 'claim_element.csv'),
              ['element_id', 'patent_id', 'claim_no', 'element_no', 'element_text', 'element_smiles'], ce_rows)

    gg_rows = build_grounding_gold(args.citations_dir, known_patent_ids)
    write_csv(os.path.join(out_dir, 'grounding_gold.csv'),
              ['ground_id', 'query_patent_id', 'claim_element_id', 'target_rel_claims',
               'cited_patent_id', 'evidence_type', 'evidence_unit_id', 'source'], gg_rows)

    print('\nDone. exam_case.csv / case_relevance.csv (Task 3) remain empty because '
          'original examination-outcome/case data could not be found locally.')


if __name__ == '__main__':
    main()
