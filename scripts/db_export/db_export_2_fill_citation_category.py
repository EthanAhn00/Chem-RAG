"""Fill data/db_export/examiner_citation.csv category (X/Y/A) from Project/epo_dataset/parsed_citations_*.xlsx. Run db_export_1_core_tables.py first.
Read columns: ID (target EP publication, e.g. 'EP3253046A3'); dnum (cited EP/US/JP/WO etc. document); country (retain EP, matching current cited_patent_id); category (X/Y/A and comma-separated D/P/T/E/L flags, stored unchanged).
Do not fill rejection_type (102/103): USPTO statutory grounds do not map one-to-one to EPO categories; alignment was undecided at the 8/24 meeting. For repeated target/cited rows choose severity X>Y>A.
Requires openpyxl. Run python3 scripts/db_export/db_export_2_fill_citation_category.py [--citations-dir /path/to/epo_dataset].
"""
import os
import sys
import csv
import argparse
import glob

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from config.settings import DATA_DIR, PROJECT_DIR

try:
    import openpyxl
except ImportError:
    print('[FAIL] openpyxl is not installed: pip install openpyxl')
    sys.exit(1)

SEVERITY = {'X': 0, 'Y': 1, 'A': 2}


def severity_rank(category):
    primary = (category or '').split(',')[0].strip()
    return SEVERITY.get(primary, 99)


def build_category_map(citations_dir):
    xlsx_files = sorted(glob.glob(os.path.join(citations_dir, 'parsed_citations_*.xlsx')))
    if not xlsx_files:
        print(f'[FAIL] {citations_dir} contains no parsed_citations_*.xlsx files.')
        sys.exit(1)

    category_map = {}
    ep_row_count, total_rows = 0, 0

    for path in xlsx_files:
        print(f'Reading: {os.path.basename(path)} ...')
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        idx = {h: i for i, h in enumerate(header)}
        required = ('ID', 'dnum', 'country', 'category')
        if not all(c in idx for c in required):
            print(f'  [WARNING] {path}: required columns missing; skipping. Actual columns: {header}')
            wb.close()
            continue

        for row in rows:
            total_rows += 1
            if row[idx['country']] != 'EP':
                continue
            target_id, cited_id, category = row[idx['ID']], row[idx['dnum']], row[idx['category']]
            if not target_id or not cited_id or not category:
                continue
            ep_row_count += 1
            key = (target_id, cited_id)
            existing = category_map.get(key)
            if existing is None or severity_rank(category) < severity_rank(existing):
                category_map[key] = category
        wb.close()

    print(f'\nScanned {len(xlsx_files)} annual files, {total_rows} rows.')
    print(f'EP citation rows: {ep_row_count} -> unique (target,cited) pairs: {len(category_map)}.')
    return category_map


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--citations-dir', default=os.path.join(os.path.dirname(PROJECT_DIR), 'epo_dataset'))
    args = parser.parse_args()

    ec_path = os.path.join(DATA_DIR, 'db_export', 'examiner_citation.csv')
    if not os.path.exists(ec_path):
        print(f'[FAIL] {ec_path} is missing. Run python3 scripts/db_export/db_export_1_core_tables.py first.')
        sys.exit(1)

    category_map = build_category_map(args.citations_dir)

    with open(ec_path, 'r', encoding='utf-8', newline='') as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames
        rows = list(reader)

    filled = 0
    category_value_counts = {}
    for row in rows:
        category = category_map.get((row['query_patent_id'], row['cited_patent_id']))
        if category:
            row['category'] = category
            filled += 1
            category_value_counts[category] = category_value_counts.get(category, 0) + 1

    with open(ec_path, 'w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f'\n{ec_path} updated: total {len(rows)}; populated {filled} categories '
          f'({filled / len(rows) * 100:.1f}%).')
    print('Distribution of populated categories:', category_value_counts)
    unmatched = len(rows) - filled
    if unmatched:
        print(f'\n[NOTE] {unmatched} unmatched records: possibly outside 2012~2023 or differing kind codes.')


if __name__ == '__main__':
    main()
